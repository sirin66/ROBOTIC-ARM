import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Helper styling function
    def set_font(run, name='Calibri', size=11, bold=False, italic=False, color=None):
        run.font.name = name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Micro Project: 3-DOF Robotic Manipulator")
    set_font(run_title, name='Arial', size=20, bold=True, color=RGBColor(26, 54, 93))

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run("Kinematic Modeling, Forward Kinematics Formulation & 3D Simulation")
    set_font(run_sub, name='Calibri', size=12, italic=True, color=RGBColor(74, 85, 104))

    # Meta / Student Info Table
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    data = [
        ("Project Title:", "3-DOF Revolute Robotic Manipulator (RRR)"),
        ("Student Name:", "SIRIN J DEVASSIA"),
        ("Roll No:", "28"),
        ("Department:", "Robotics and Artificial Intelligence (RAI)"),
        ("Academic Year:", "2025 - 2026")
    ]

    for idx, (label, val) in enumerate(data):
        row = info_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.5)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(3)
        p_lbl.paragraph_format.space_after = Pt(3)
        r_lbl = p_lbl.add_run(label)
        set_font(r_lbl, name='Calibri', size=10.5, bold=True, color=RGBColor(45, 55, 72))

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(3)
        p_val.paragraph_format.space_after = Pt(3)
        r_val = p_val.add_run(val)
        set_font(r_val, name='Calibri', size=10.5, bold=False, color=RGBColor(45, 55, 72))

        # Shading
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{"F7FAFC" if idx % 2 == 0 else "EDF2F7"}"/>')
        cell_lbl._tc.get_or_add_tcPr().append(shading)
        shading2 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{"F7FAFC" if idx % 2 == 0 else "EDF2F7"}"/>')
        cell_val._tc.get_or_add_tcPr().append(shading2)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_heading(num_str, title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r_num = h.add_run(num_str + " ")
        set_font(r_num, name='Arial', size=13, bold=True, color=RGBColor(43, 108, 176))
        r_title = h.add_run(title_str)
        set_font(r_title, name='Arial', size=13, bold=True, color=RGBColor(26, 32, 44))

        # Underline accent line
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(8)
        p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="3182CE"/></w:pBdr>')
        p_line._p.get_or_add_pPr().append(p_line_border)

    # ----------------------------------------------------
    # SECTION 1: ALGORITHM (PSEUDO-CODE)
    # ----------------------------------------------------
    add_section_heading("1.", "Algorithm (Pseudo-code)")

    algo_desc = doc.add_paragraph()
    algo_desc.paragraph_format.space_after = Pt(6)
    r_ad = algo_desc.add_run("The following step-by-step logic defines the geometric initialization, kinematic transformation pipeline, and 3D simulation of the 3-DOF articulated manipulator:")
    set_font(r_ad, name='Calibri', size=11, color=RGBColor(45, 55, 72))

    algo_steps = [
        ("Step 1: Link Parameter Definition", "Define the physical link dimensions: Base height L1, Upper arm length L2, and Forearm length L3 in meters."),
        ("Step 2: DH Parameter Formulation", "Assign Standard Denavit-Hartenberg parameters for each link:\n• Link 1: d = L1, a = 0, alpha = +pi/2\n• Link 2: d = 0, a = L2, alpha = 0\n• Link 3: d = 0, a = L3, alpha = 0"),
        ("Step 3: Robot Assembly", "Instantiate RevoluteDH joint objects and assemble the serial-link DHRobot model."),
        ("Step 4: Input Joint Angles", "Acquire joint variables q1, q2, and q3 in degrees and convert them into radians."),
        ("Step 5: Homogeneous Transformation Calculation", "Compute the individual link transformation matrices A1, A2, and A3, and evaluate the composite end-effector transformation matrix:\n    T_0_3 = A1 * A2 * A3"),
        ("Step 6: End-Effector Position Extraction", "Extract the translational coordinates (X, Y, Z) from the fourth column of matrix T_0_3:\n    X = cos(q1) * [L2 * cos(q2) + L3 * cos(q2 + q3)]\n    Y = sin(q1) * [L2 * cos(q2) + L3 * cos(q2 + q3)]\n    Z = L1 + L2 * sin(q2) + L3 * sin(q2 + q3)"),
        ("Step 7: 3D Kinematic Plotting", "Pass the joint angles to the robotics visualization environment to render the links, joint axes, and coordinate frames in 3D Euclidean space, and save the resulting graphical plot to file.")
    ]

    for title, desc in algo_steps:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r_st = p.add_run(title + "\n")
        set_font(r_st, name='Calibri', size=10.5, bold=True, color=RGBColor(43, 108, 176))
        r_sd = p.add_run(desc)
        set_font(r_sd, name='Calibri', size=10, bold=False, color=RGBColor(74, 85, 104))

    # ----------------------------------------------------
    # SECTION 2: PROGRAM
    # ----------------------------------------------------
    add_section_heading("2.", "Program")

    p_prog_desc = doc.add_paragraph()
    p_prog_desc.paragraph_format.space_after = Pt(6)
    r_pd = p_prog_desc.add_run("Core Python implementation utilizing Peter Corke's Robotics Toolbox for forward kinematics calculation and 3D simulation:")
    set_font(r_pd, name='Calibri', size=11, color=RGBColor(45, 55, 72))

    code_content = """import numpy as np
import roboticstoolbox as rtb
import matplotlib.pyplot as plt

class Manipulator3DOF:
    def __init__(self, l1=0.5, l2=0.4, l3=0.3):
        self.l1 = l1
        self.l2 = l2
        self.l3 = l3
        
        self.link1 = rtb.RevoluteDH(d=self.l1, a=0.0, alpha=np.pi/2)
        self.link2 = rtb.RevoluteDH(d=0.0, a=self.l2, alpha=0.0)
        self.link3 = rtb.RevoluteDH(d=0.0, a=self.l3, alpha=0.0)
        
        self.robot = rtb.DHRobot([self.link1, self.link2, self.link3], name="3DOF_Manipulator")

    def forward_kinematics(self, q):
        return self.robot.fkine(q)

    def plot_arm(self, q, save_path=None):
        self.robot.plot(q, backend="pyplot", block=False)
        if save_path:
            plt.savefig(save_path, dpi=300)
        plt.close()

if __name__ == "__main__":
    arm = Manipulator3DOF(l1=0.5, l2=0.4, l3=0.3)
    
    test_angles = [
        [0.0, 0.0, 0.0],
        [30.0, 45.0, 30.0],
        [-45.0, 60.0, -30.0]
    ]
    
    for i, angles in enumerate(test_angles, 1):
        q_rad = np.deg2rad(angles)
        T = arm.forward_kinematics(q_rad)
        pos = T.t
        
        print(f"Test Configuration {i}:")
        print(f"Joint Angles (deg): q1 = {angles[0]}, q2 = {angles[1]}, q3 = {angles[2]}")
        print("Homogeneous Transformation Matrix:")
        print(np.round(T.A, 4))
        print(f"End-Effector Coordinates (m): X = {pos[0]:.4f}, Y = {pos[1]:.4f}, Z = {pos[2]:.4f}\\n")
    
    arm.plot_arm(np.deg2rad([30.0, 45.0, 30.0]), save_path="manipulator_3d_pose.png")"""

    code_table = doc.add_table(rows=1, cols=1)
    code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    code_cell = code_table.rows[0].cells[0]
    code_cell.width = Inches(6.5)

    shading_c = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    code_cell._tc.get_or_add_tcPr().append(shading_c)
    border_c = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="3182CE"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    code_cell._tc.get_or_add_tcPr().append(border_c)

    cp = code_cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.line_spacing = 1.15
    rc = cp.add_run(code_content)
    set_font(rc, name='Consolas', size=9, color=RGBColor(26, 32, 44))

    # ----------------------------------------------------
    # SECTION 3: RESULT
    # ----------------------------------------------------
    add_section_heading("3.", "Result")

    r_desc = doc.add_paragraph()
    r_desc.paragraph_format.space_after = Pt(6)
    r_rd = r_desc.add_run("The script was executed for three distinct joint angle test configurations to evaluate the spatial displacement of the end-effector:")
    set_font(r_rd, name='Calibri', size=11, color=RGBColor(45, 55, 72))

    console_output = """Test Configuration 1:
Joint Angles (deg): q1 = 0.0, q2 = 0.0, q3 = 0.0
Homogeneous Transformation Matrix:
[[ 1.  0.  0.  0.7]
 [ 0.  0. -1.  0. ]
 [ 0.  1.  0.  0.5]
 [ 0.  0.  0.  1. ]]
End-Effector Coordinates (m): X = 0.7000, Y = 0.0000, Z = 0.5000

Test Configuration 2:
Joint Angles (deg): q1 = 30.0, q2 = 45.0, q3 = 30.0
Homogeneous Transformation Matrix:
[[ 0.2241 -0.8365  0.5     0.3122]
 [ 0.1294 -0.483  -0.866   0.1802]
 [ 0.9659  0.2588  0.      1.0726]
 [ 0.      0.      0.      1.    ]]
End-Effector Coordinates (m): X = 0.3122, Y = 0.1802, Z = 1.0726

Test Configuration 3:
Joint Angles (deg): q1 = -45.0, q2 = 60.0, q3 = -30.0
Homogeneous Transformation Matrix:
[[ 0.6124 -0.3536 -0.7071  0.3251]
 [-0.6124  0.3536 -0.7071 -0.3251]
 [ 0.5     0.866   0.      0.9964]
 [ 0.      0.      0.      1.    ]]
End-Effector Coordinates (m): X = 0.3251, Y = -0.3251, Z = 0.9964"""

    out_table = doc.add_table(rows=1, cols=1)
    out_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    out_cell = out_table.rows[0].cells[0]
    out_cell.width = Inches(6.5)

    shading_o = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
    out_cell._tc.get_or_add_tcPr().append(shading_o)

    op = out_cell.paragraphs[0]
    op.paragraph_format.space_before = Pt(6)
    op.paragraph_format.space_after = Pt(6)
    op.paragraph_format.line_spacing = 1.15
    ro = op.add_run(console_output)
    set_font(ro, name='Consolas', size=9, color=RGBColor(226, 232, 240))

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Summary Table of Results
    p_tab_title = doc.add_paragraph()
    p_tab_title.paragraph_format.space_before = Pt(8)
    p_tab_title.paragraph_format.space_after = Pt(4)
    r_tt = p_tab_title.add_run("Table 1: Kinematic Comparison Across Configurations")
    set_font(r_tt, name='Calibri', size=11, bold=True, color=RGBColor(45, 55, 72))

    res_table = doc.add_table(rows=4, cols=5)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_table.autofit = False

    headers = ["Config", "Joint Angles (q1, q2, q3)", "X (m)", "Y (m)", "Z (m)"]
    row_data = [
        ["1 (Home)", "[0.0°, 0.0°, 0.0°]", "0.7000", "0.0000", "0.5000"],
        ["2 (Reach)", "[30.0°, 45.0°, 30.0°]", "0.3122", "0.1802", "1.0726"],
        ["3 (Pick)", "[-45.0°, 60.0°, -30.0°]", "0.3251", "-0.3251", "0.9964"]
    ]

    col_widths = [Inches(1.0), Inches(2.2), Inches(1.1), Inches(1.1), Inches(1.1)]

    for c_idx, h_text in enumerate(headers):
        cell = res_table.rows[0].cells[c_idx]
        cell.width = col_widths[c_idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        set_font(r, name='Calibri', size=10, bold=True, color=RGBColor(255, 255, 255))
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="2B6CB0"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    for r_idx, row in enumerate(row_data, 1):
        for c_idx, val in enumerate(row):
            cell = res_table.rows[r_idx].cells[c_idx]
            cell.width = col_widths[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            set_font(r, name='Calibri', size=10, bold=(c_idx == 0), color=RGBColor(45, 55, 72))
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{"F7FAFC" if r_idx % 2 == 1 else "EDF2F7"}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3D Plot Embedding
    p_img_title = doc.add_paragraph()
    p_img_title.paragraph_format.space_before = Pt(10)
    p_img_title.paragraph_format.space_after = Pt(6)
    r_it = p_img_title.add_run("3D Simulation Plot of the Manipulator (Configuration 2):")
    set_font(r_it, name='Calibri', size=11, bold=True, color=RGBColor(45, 55, 72))

    image_path = "manipulator_3d_pose.png"
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        doc.add_picture(image_path, width=Inches(5.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(16)
        r_cap = p_cap.add_run("Figure 1: Generated 3D Robotic Manipulator Model at Joint Angles q = [30.0°, 45.0°, 30.0°]")
        set_font(r_cap, name='Calibri', size=9.5, italic=True, color=RGBColor(100, 116, 139))

    output_filename = "Micro_Project_3DOF_Robotic_Manipulator.docx"
    doc.save(output_filename)
    print(f"[OK] Successfully generated report: {output_filename}")

if __name__ == "__main__":
    create_report()
